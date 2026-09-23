"""Genera las hojas individuales de la corrección del 23-09-2026."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).parent
DOCS = ROOT / "documentos"
SOURCES = ROOT / "fuentes-md"
INK = RGBColor(31, 38, 40)
MUTED = RGBColor(80, 96, 99)
ACCENT = RGBColor(19, 104, 101)

PARTICIPANT_BLOCKS = [
    ("field", "Código: __________________    Día 1: ________________    Día 2: ________________"),
    ("body", "Elige una actividad, una aplicación y un tiempo. Deja el parlante cerca de lo que necesitas para comenzar. Si suena, tú decides qué hacer. Si algo falla o molesta, detén la prueba y avisa al responsable."),
    ("section", "Día 1"),
    ("check", "¿Activaste Relevo?  ☐ Sí   ☐ No       ¿Sonó el aviso?  ☐ Sí   ☐ No   ☐ No sé"),
    ("field", "¿Qué querías hacer y dónde dejaste el parlante? ______________________________"),
    ("field", "__________________________________________________________________________"),
    ("field", "¿Qué hiciste después? Si no lo activaste, ¿qué pasó? ____________________________"),
    ("field", "__________________________________________________________________________"),
    ("section", "Día 2"),
    ("check", "¿Activaste Relevo?  ☐ Sí   ☐ No       ¿Sonó el aviso?  ☐ Sí   ☐ No   ☐ No sé"),
    ("field", "¿Qué querías hacer y dónde dejaste el parlante? ______________________________"),
    ("field", "__________________________________________________________________________"),
    ("field", "¿Qué hiciste después? Si no lo activaste, ¿qué pasó? ____________________________"),
    ("field", "__________________________________________________________________________"),
    ("section", "Después de los dos días"),
    ("body", "Marca un número por línea. En preparación y claridad: 1 = muy difícil, 7 = muy fácil. En sonido: 1 = muy incómodo, 7 = muy cómodo."),
    ("check", "Preparar Relevo                 ①   ②   ③   ④   ⑤   ⑥   ⑦"),
    ("check", "Entender el aviso               ①   ②   ③   ④   ⑤   ⑥   ⑦"),
    ("check", "Comodidad del sonido       ①   ②   ③   ④   ⑤   ⑥   ⑦"),
    ("field", "¿Qué cambiarías o qué te hizo dudar? __________________________________________"),
    ("field", "__________________________________________________________________________"),
    ("note", "También sirve saber que decidiste no usarlo. Contacto del responsable: ______________________________"),
]


PAGES = [
    (
        "01-prototipo-construido",
        "Prototipo construido",
        "Presentación para la corrección · Relevo · 23 de septiembre de 2026",
        [
            ("lead", "Relevo relaciona una intención elegida por la persona con un aviso sonoro situado cerca de lo que necesita para comenzar. El prototipo reúne una aplicación Android y un parlante Bluetooth portátil."),
            ("section", "Qué se puede mostrar hoy"),
            ("bullet", "APK Android 2.3 compilado y proyecto de Android Studio disponible. Requiere Android 12 o posterior."),
            ("bullet", "En la app: consentimiento inicial, tutorial, creación de actividad, elección de aplicación y tiempo, ubicación del parlante, activación, aviso y cierre."),
            ("bullet", "Registro local de sesiones y eventos; sincronización remota configurada. No se leen mensajes ni contenidos de otras aplicaciones."),
            ("section", "Recorrido de demostración"),
            ("check", "Abrir la app y mostrar consentimiento, tutorial y permisos."),
            ("check", "Crear una actividad, seleccionar una aplicación y fijar un tiempo breve."),
            ("check", "Conectar el parlante, dejarlo junto a la actividad y activar Relevo."),
            ("check", "Usar la aplicación elegida, esperar el aviso, silenciarlo y cerrar el ciclo."),
            ("section", "Estado que debe comunicarse con precisión"),
            ("body", "La versión 2.3 compiló y pasó pruebas unitarias. No se completó una inspección visual estable en emulador ni la prueba física con un teléfono y parlante concretos. La emisión exclusiva por Bluetooth, el segundo plano y la desconexión deben comprobarse antes de entregar el kit a participantes."),
            ("note", "Para la sesión: llevar el APK, el parlante, un teléfono Android compatible, esta hoja y los instrumentos. Si falla la demostración física, mostrar el flujo y declarar el fallo; no atribuirle resultados de uso."),
        ],
    ),
    (
        "02-indicadores-validacion",
        "Indicadores de validación",
        "Pauta para revisión docente · Prototipo de interacción · dos días por persona",
        [
            ("lead", "La evaluación busca saber si se entiende y controla el paso desde una intención elegida hasta un aviso situado. Cada indicador distingue lo observado de lo que la persona declara."),
            ("section", "Configuración y comprensión"),
            ("bullet", "Autonomía: completa actividad, app, tiempo y ubicación sin ayuda / con ayuda / no completa. Anotar en qué paso aparece la duda."),
            ("bullet", "Modelo de conteo: explica qué app acumula tiempo y si el conteo se pausa al salir. Registrar explicación literal, no solo sí/no."),
            ("section", "Señal y comportamiento"),
            ("bullet", "Ubicación: dónde dejó el parlante y por qué ese lugar le sirve para comenzar la actividad."),
            ("bullet", "Funcionamiento: aviso tras el tiempo acordado, salida de audio, segundo plano, silencio y desconexión. Registrar hora y fallo técnico por separado."),
            ("bullet", "Pertinencia: si el aviso llegó en un momento útil, molesto o confuso; pedir un ejemplo concreto."),
            ("section", "Respuesta y control"),
            ("bullet", "Recuperación: si recuerda la actividad y cómo pensaba comenzar sin que se le den pistas."),
            ("bullet", "Control: si puede silenciar, cambiar de idea o cerrar sin sentirse obligado a realizar la actividad."),
            ("bullet", "Comodidad: puntuación 1–7 para preparación, claridad del aviso y sonido; acompañar cada valor con una razón."),
            ("section", "Regla de lectura"),
            ("body", "Un aviso registrado no demuestra que fue oído. Elegir «comencé» es un autorreporte, no prueba de actividad física. No calcular porcentajes ni declarar eficacia causal con dos casos; comparar recorridos, fallos y explicaciones."),
            ("note", "Base metodológica: observación de tareas de usabilidad (Nielsen Norman Group, 2019) y diseño centrado en las personas (ISO 9241-210:2019). Indicadores sujetos a ajuste en la corrección."),
        ],
    ),
    (
        "03-pauta-de-testeo",
        "Pauta para dirigir el testeo",
        "Uso del investigador · una persona · dos días · sin resultados anticipados",
        [
            ("section", "Antes de invitar"),
            ("check", "Probar en el teléfono y parlante reales: permisos, audio Bluetooth, app en segundo plano, pausa del conteo, silencio y desconexión."),
            ("check", "Completar contacto del responsable; revisar consentimiento y eliminación de datos locales y remotos; registrar modelos y APK."),
            ("check", "Preparar un parlante, APK y hojas de ambos días con copia de repuesto. Guardar consentimiento firmado aparte."),
            ("section", "Día 1 · preparación acompañada"),
            ("body", "Tras explicar y obtener consentimiento, pedir sin indicar botones: «Piensa en algo que quieras empezar cuando sueles seguir usando otra aplicación. Prepara Relevo para que te avise al cumplirse el tiempo que elijas». Observar dudas, ayudas y dónde deja el parlante. Antes del aviso: «¿Qué esperas que ocurra?»"),
            ("body", "Después del aviso: «¿Qué te recordó?» y «¿Cómo pensabas empezar?». Pedir que cierre el relevo como prefiera. Anotar cualquier ayuda entregada."),
            ("section", "Día 2 · uso en contexto y cierre"),
            ("body", "La persona decide si y cuándo usarlo. Un día sin activación también es un dato. Al cierre, revisar su hoja sin corregir sus respuestas, preguntar por un momento útil y uno inoportuno, recuperar el parlante y registrar problemas técnicos aparte."),
            ("section", "Límites durante la prueba"),
            ("bullet", "No revisar mensajes, búsquedas ni contenidos de otras aplicaciones. No solicitar capturas o grabaciones sin consentimiento adicional."),
            ("bullet", "Ante sonido que no se puede detener, exposición de datos o aviso fuera de lo acordado: detener la prueba y documentar el incidente."),
            ("note", "Registro: separar hecho observado, palabras textuales, dato de la app e interpretación. No presentar esta pauta como un testeo ya realizado."),
        ],
    ),
    (
        "04-consentimiento",
        "Consentimiento informado",
        "Borrador para revisión docente · Universidad Diego Portales · responsable: Johan Yantén",
        [
            ("body", "Se me invita a probar durante dos días Relevo: una aplicación Android y un parlante Bluetooth. Puedo elegir una actividad, una aplicación cuyo tiempo de uso observar y un límite. Si se cumple, el sistema intenta emitir un aviso. Se me pedirá una hoja de registro y una conversación breve al final."),
            ("section", "Mi decisión y posibles molestias"),
            ("body", "Participar es voluntario. Puedo no responder preguntas, silenciar el parlante o retirarme sin dar razones ni consecuencias. El sonido puede molestar o aparecer en un momento inoportuno; puedo bajar el volumen o detener la prueba. No debo entregar claves ni información de salud."),
            ("section", "Datos y privacidad"),
            ("body", "La app solicita acceso al tiempo de uso de aplicaciones y permiso de notificaciones. Guarda un código aleatorio, actividad, primer paso, lugar declarado, app elegida, límite, tiempo acumulado, momentos de activación, aviso y cierre, y mi respuesta final solo si la doy. Los datos quedan en el teléfono y pueden sincronizarse con una base remota. No se registran mensajes, búsquedas, fotos ni contenido de pantalla. La hoja y esta firma se guardan por separado."),
            ("body", "Los resultados académicos se presentarán sin mi nombre. Puedo solicitar la eliminación de mis registros mediante el código entregado. Se propone eliminarlos a más tardar el 13 de enero de 2027. La posibilidad de ubicar y borrar los registros locales y remotos debe comprobarse antes de comenzar la prueba."),
            ("note", "Este consentimiento no reemplaza los permisos de Android ni la aceptación que muestra la app. Contacto para preguntas o retiro: ____________________________________"),
            ("section", "Aceptación"),
            ("body", "He leído o escuchado la información, pude preguntar y acepto participar durante dos días."),
            ("field", "Código: _______________     Fecha: __________________________"),
            ("field", "Nombre y firma de participante: ______________________________________________"),
            ("field", "Firma del responsable: ______________________________________________________"),
        ],
    ),
    (
        "05-registro-participante",
        "Registro de uso · dos días",
        "Hoja de participante · completa ambos días y entrégala al finalizar la prueba",
        PARTICIPANT_BLOCKS,
    ),
    (
        "06-registro-repuesto",
        "Registro de uso · dos días",
        "Copia de repuesto · completa ambos días y entrégala al finalizar la prueba",
        PARTICIPANT_BLOCKS,
    ),
    (
        "07-plan-de-trabajo",
        "Plan de trabajo individual",
        "Johan Yantén · resumen para la corrección · septiembre–diciembre de 2026",
        [
            ("lead", "Objetivo: llegar a noviembre con producto y memoria integrados. Diciembre queda reservado principalmente para revisar, ensayar la defensa y resolver errores menores."),
            ("section", "Esta semana · 23 al 30 de septiembre"),
            ("check", "Mostrar prototipo e instrumentos en la corrección; registrar observaciones y decidir cambios."),
            ("check", "Comprobar con equipo real audio, segundo plano, permisos y desconexión; cerrar contacto y eliminación de datos."),
            ("section", "Octubre · validación e iteración"),
            ("check", "Realizar piloto técnico y prueba de dos días con participantes consentidos, solo cuando el kit esté verificado."),
            ("check", "Comparar observaciones, registros y relatos; corregir app, señal, ubicación y forma física donde corresponda."),
            ("check", "Documentar costos, montaje, fallos, accesibilidad y decisiones de diseño."),
            ("section", "Noviembre · integración y cierre"),
            ("check", "Cerrar prototipo integrado y comprobar que se instala, funciona y puede demostrarse."),
            ("check", "Integrar resultados reales y límites en la memoria; revisar citas, figuras, anexos y coherencia del producto."),
            ("section", "Diciembre · entrega y defensa"),
            ("check", "Verificar archivos y respaldos; ensayar exposición y preguntas críticas; corregir solo detalles menores."),
            ("note", "Plan detallado en la pestaña «Johan» del archivo de planificación 2026. Fechas y prioridades pueden ajustarse tras la única corrección disponible en esta sesión."),
        ],
    ),
    (
        "08-pendientes-criticos",
        "Pendientes antes de probar",
        "Lista de cierre técnico y ético · no equivale a resultados de validación",
        [
            ("lead", "El prototipo y los instrumentos están preparados para revisión docente. La entrega de un kit a participantes requiere completar estas comprobaciones."),
            ("section", "Kit físico y aplicación"),
            ("check", "Registrar modelo de teléfono Android, parlante, versión APK y configuración Bluetooth."),
            ("check", "Verificar sonido solo en el parlante, con la app fuera de pantalla, y que se detiene al silenciar."),
            ("check", "Confirmar que el conteo sigue solo en la aplicación elegida y se pausa al salir."),
            ("check", "Probar desconexión, reconexión, volumen, batería, permisos y cierre de ciclo."),
            ("section", "Participación y datos"),
            ("check", "Completar un medio de contacto real en el consentimiento y las hojas."),
            ("check", "Alinear consentimiento en papel y app: finalidad, datos, plazo, retiro y eliminación."),
            ("check", "Mostrar un código a cada participante y demostrar localización y borrado local/remoto de sus datos."),
            ("check", "Acordar los instrumentos con el docente antes de reclutar o recoger datos."),
            ("section", "Después de la corrección"),
            ("check", "Registrar cambios pedidos, responsable y fecha; actualizar el plan y la versión de los instrumentos."),
            ("note", "Si una condición crítica falla, detener esa parte del ensayo y corregirla antes de atribuir dificultades a la persona. Hoy no se reportan testeos con usuarios."),
        ],
    ),
]


def add_text(doc, kind, content):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(5 if kind in {"body", "lead"} else 3)
    para.paragraph_format.line_spacing = 1.12
    if kind == "section":
        para.paragraph_format.space_before = Pt(9)
        run = para.add_run(content)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = ACCENT
    elif kind == "bullet":
        para.style = "List Bullet"
        run = para.add_run(content)
        run.font.size = Pt(10.5)
    elif kind == "check":
        if "☐" not in content:
            content = "☐  " + content
        for index, piece in enumerate(content.split("☐")):
            if index:
                box = para.add_run("☐")
                box.font.name = "Segoe UI Symbol"
                box.font.size = Pt(16)
            if piece:
                part = para.add_run(piece)
                part.font.size = Pt(10.5)
    else:
        run = para.add_run(content)
        run.font.size = Pt(10.5 if kind != "lead" else 11)
        if kind == "lead":
            run.bold = True
        if kind == "note":
            run.font.color.rgb = MUTED
    for run in para.runs:
        run.font.name = run.font.name or "Arial"
        if run.font.color.rgb is None:
            run.font.color.rgb = INK


def build_docx(name, title, subtitle, blocks):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3)
    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Cm(0.55)
    bullet.paragraph_format.first_line_indent = Cm(-0.25)
    eyebrow = doc.add_paragraph("RELEVO  /  ENTREGA DE CORRECCIÓN")
    eyebrow.runs[0].font.name = "Arial"
    eyebrow.runs[0].font.size = Pt(10)
    eyebrow.runs[0].font.bold = True
    eyebrow.runs[0].font.color.rgb = ACCENT
    eyebrow.paragraph_format.space_after = Pt(8)
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = INK
    sub = doc.add_paragraph(subtitle)
    sub.runs[0].font.name = "Arial"
    sub.runs[0].font.size = Pt(10.5)
    sub.runs[0].font.color.rgb = MUTED
    sub.paragraph_format.space_after = Pt(12)
    for kind, content in blocks:
        add_text(doc, kind, content)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot = footer.add_run("RELEVO · 23.09.2026 · " + name[:2])
    foot.font.name = "Arial"
    foot.font.size = Pt(10)
    foot.font.color.rgb = MUTED
    doc.save(DOCS / f"{name}.docx")


def build_markdown(name, title, subtitle, blocks):
    lines = [f"# {title}", "", subtitle, ""]
    for kind, content in blocks:
        if kind == "section":
            lines.extend([f"## {content}", ""])
        elif kind in {"bullet", "check"}:
            lines.append(f"- {content}")
        else:
            lines.extend([content, ""])
    lines.extend([
        "## Registro de cambios y alcance",
        "",
        "- **Qué se cambió:** esta hoja se separó del paquete único y se actualizó para el prototipo Android 2.3.",
        "- **Cómo era antes:** pauta, consentimiento, indicadores, registro y pendientes compartían un solo archivo extenso referido a la versión 2.2.",
        "- **Por qué:** permitir lectura y uso independientes, campos amplios y una hoja imprimible por documento.",
        "- **Alcance:** material listo para corrección docente; no acredita ensayos físicos ni resultados con participantes.",
        "",
    ])
    (SOURCES / f"{name}.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    DOCS.mkdir(exist_ok=True)
    SOURCES.mkdir(exist_ok=True)
    for page in PAGES:
        build_markdown(*page)
        build_docx(*page)


if __name__ == "__main__":
    main()
